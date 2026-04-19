"""
AI Resume Coach — FastAPI Backend
Provides: resume analysis, scoring, JD matching, mock interview evaluation,
          history tracking, and downloadable PDF report generation.
"""

from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from openai import OpenAI
import os, re, io, json
from datetime import datetime
from dotenv import load_dotenv
import PyPDF2
from docx import Document
from fpdf import FPDF

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
load_dotenv()
api_key = os.getenv("OPENAI_API_KEY", "")
PLACEHOLDER_KEYS = {"", "your_openai_api_key_here", "sk-xxx", "your-key-here"}
use_ai = api_key not in PLACEHOLDER_KEYS and api_key.startswith("sk-")
client = OpenAI(api_key=api_key) if use_ai else None

app = FastAPI(title="AI Resume Coach API")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

# Register builder module
from builder import router as builder_router
app.include_router(builder_router)

HISTORY_FILE = "history.json"

# ---------------------------------------------------------------------------
# Pydantic Models
# ---------------------------------------------------------------------------
class ScoreBreakdown(BaseModel):
    ats_score: int
    readability_score: int
    impact_score: int
    overall_score: int

class AnalysisResult(BaseModel):
    suggestions: list[str]
    interview_questions: list[str]
    scores: ScoreBreakdown
    detected_skills: list[str]
    missing_sections: list[str]
    resume_text: str  # returned so frontend can reuse it

class JDMatchRequest(BaseModel):
    resume_text: str
    job_description: str

class JDMatchResult(BaseModel):
    match_score: int
    matched_keywords: list[str]
    missing_keywords: list[str]
    skill_gaps: list[dict]

class InterviewEvalRequest(BaseModel):
    question: str
    answer: str
    resume_context: str = ""

class InterviewEvalResult(BaseModel):
    confidence_score: int
    clarity_score: int
    relevance_score: int
    overall_score: int
    feedback: str
    improvement_tips: list[str]

class ReportRequest(BaseModel):
    filename: str = "resume"
    suggestions: list[str] = []
    interview_questions: list[str] = []
    scores: dict = {}
    detected_skills: list[str] = []
    missing_sections: list[str] = []
    jd_match: dict | None = None

# ---------------------------------------------------------------------------
# Text Extraction
# ---------------------------------------------------------------------------
def extract_text(file_bytes: bytes, filename: str) -> str:
    text = ""
    if filename.endswith(".pdf"):
        reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
        for page in reader.pages:
            text += page.extract_text() or ""
    elif filename.endswith(".docx"):
        doc = Document(io.BytesIO(file_bytes))
        for para in doc.paragraphs:
            text += para.text + "\n"
    return text

# ---------------------------------------------------------------------------
# Skill Detection
# ---------------------------------------------------------------------------
TECH_KEYWORDS = {
    "python": "Python", "java ": "Java", "javascript": "JavaScript",
    "typescript": "TypeScript", "react": "React", "angular": "Angular",
    "vue": "Vue.js", "node": "Node.js", "express": "Express.js",
    "sql": "SQL", "nosql": "NoSQL", "mongodb": "MongoDB",
    "postgresql": "PostgreSQL", "mysql": "MySQL", "redis": "Redis",
    "machine learning": "Machine Learning", "deep learning": "Deep Learning",
    "artificial intelligence": "AI", "nlp": "NLP",
    "computer vision": "Computer Vision", "aws": "AWS", "azure": "Azure",
    "gcp": "GCP", "docker": "Docker", "kubernetes": "Kubernetes",
    "terraform": "Terraform", "jenkins": "Jenkins", "git": "Git",
    "ci/cd": "CI/CD", "rest api": "REST APIs", "api": "APIs",
    "tensorflow": "TensorFlow", "pytorch": "PyTorch", "keras": "Keras",
    "pandas": "Pandas", "numpy": "NumPy", "scikit": "Scikit-learn",
    "data analysis": "Data Analysis", "data science": "Data Science",
    "agile": "Agile", "scrum": "Scrum", "html": "HTML", "css": "CSS",
    "c++": "C++", "c#": "C#", "golang": "Go", "rust": "Rust",
    "flask": "Flask", "django": "Django", "fastapi": "FastAPI",
    "spring": "Spring Boot", "linux": "Linux", "bash": "Bash",
    "tableau": "Tableau", "power bi": "Power BI", "excel": "Excel",
    "figma": "Figma", "blockchain": "Blockchain",
    "leadership": "Leadership", "communication": "Communication",
    "problem solving": "Problem Solving", "teamwork": "Teamwork",
    "project management": "Project Management",
}

def detect_skills(text: str) -> list[str]:
    t = text.lower()
    return sorted(set(name for kw, name in TECH_KEYWORDS.items() if kw in t))

# ---------------------------------------------------------------------------
# Section Detection
# ---------------------------------------------------------------------------
SECTION_MAP = {
    "Education": ["education", "academic", "university", "degree", "bachelor",
                  "master", "b.tech", "m.tech", "cgpa", "gpa"],
    "Experience": ["experience", "employment", "work history", "professional experience"],
    "Skills": ["skills", "technical skills", "technologies", "competencies"],
    "Projects": ["project", "projects"],
    "Certifications": ["certification", "certifications", "certified", "certificate"],
    "Summary/Objective": ["summary", "objective", "profile", "about me"],
}

def detect_missing_sections(text: str) -> list[str]:
    t = text.lower()
    return [sec for sec, kws in SECTION_MAP.items() if not any(k in t for k in kws)]

# ---------------------------------------------------------------------------
# Scoring Functions
# ---------------------------------------------------------------------------
def calc_ats(text: str, skills: list, missing: list) -> int:
    s = 100
    s -= len(missing) * 8
    s += min(len(skills) * 3, 20)
    if not re.search(r'[\w.+-]+@[\w-]+\.[\w.-]+', text): s -= 10
    if not re.search(r'[\+]?[\d\s\-\(\)]{7,15}', text): s -= 5
    if "linkedin" not in text.lower(): s -= 5
    return max(0, min(100, s))

def calc_readability(text: str) -> int:
    sents = [x.strip() for x in re.split(r'[.!?]+', text) if x.strip()]
    if not sents: return 50
    avg = sum(len(x.split()) for x in sents) / len(sents)
    s = 100
    if avg > 25: s -= (avg - 25) * 2
    elif avg < 8: s -= (8 - avg) * 3
    bullets = text.count('•') + text.count('- ') + text.count('* ')
    s += min(bullets * 2, 15)
    return max(0, min(100, int(s)))

def calc_impact(text: str) -> int:
    t = text.lower()
    verbs = ["achieved","improved","increased","reduced","developed","designed",
             "implemented","managed","led","created","optimized","launched",
             "built","delivered","streamlined","automated","analyzed","resolved"]
    s = 50 + min(sum(1 for v in verbs if v in t) * 4, 25)
    metrics = re.findall(r'\d+\s*(%|\+|users|customers|projects|clients|team|people)', t)
    s += min(len(metrics) * 5, 25)
    return max(0, min(100, s))

# ---------------------------------------------------------------------------
# Suggestion Generation
# ---------------------------------------------------------------------------
def generate_suggestions(text, scores, skills, missing):
    tips, t = [], text.lower()
    wc = len(text.split())
    if wc < 150:
        tips.append(f"📝 Resume is short (~{wc} words). Aim for 400–700 words.")
    elif wc > 1000:
        tips.append(f"📝 Resume is lengthy (~{wc} words). Trim to 1–2 pages.")
    if not re.search(r'[\w.+-]+@[\w-]+\.[\w.-]+', text):
        tips.append("⚠️ No email found — add it at the top.")
    if "linkedin" not in t:
        tips.append("💼 Add your LinkedIn profile URL.")
    for sec in missing:
        tips.append(f"📋 Missing '{sec}' section — consider adding it.")
    if scores.ats_score < 60:
        tips.append("🤖 Low ATS score — use standard headings and avoid images/tables.")
    if scores.impact_score < 60:
        tips.append("🎯 Add quantifiable achievements (e.g., 'Reduced load time by 40%').")
    if scores.readability_score < 60:
        tips.append("📖 Use concise bullets instead of long paragraphs.")
    verbs = ["achieved","improved","optimized","streamlined","automated"]
    if sum(1 for v in verbs if v in t) < 2:
        tips.append("💪 Use stronger action verbs (achieved, optimized, streamlined).")
    if len(skills) < 3:
        tips.append("🛠️ Add a dedicated 'Skills' section with your tech stack.")
    return tips[:8]

# ---------------------------------------------------------------------------
# Interview Question Generation
# ---------------------------------------------------------------------------
def generate_questions(text, skills):
    t, qs = text.lower(), []
    if skills:
        qs.append(f"You mention {', '.join(skills[:3])}. Walk me through a challenging project using these.")
    qs.append("Tell me about yourself and what makes you a strong candidate.")
    if any(k in t for k in ["intern","experience","worked","company"]):
        qs.append("Describe a significant challenge during your professional experience and your solution.")
    if any(k in t for k in ["project","developed","built","implemented"]):
        qs.append("Pick a project — explain the problem, approach, tech stack, and outcome.")
    if any(k in t for k in ["team","led","managed","leadership"]):
        qs.append("Describe a leadership moment or team conflict you resolved.")
    if any(k in t for k in ["machine learning","deep learning","model","ai"]):
        qs.append("Walk me through an ML model you built — algorithm choice, training, evaluation.")
    qs.append("Where do you see yourself in 3–5 years?")
    seen, uniq = set(), []
    for q in qs:
        if q not in seen: seen.add(q); uniq.append(q)
    return uniq[:6]

# ---------------------------------------------------------------------------
# Local Analysis (no API key required)
# ---------------------------------------------------------------------------
def local_analyze(text: str, filename: str = "resume") -> AnalysisResult:
    skills = detect_skills(text)
    missing = detect_missing_sections(text)
    ats = calc_ats(text, skills, missing)
    read = calc_readability(text)
    imp = calc_impact(text)
    overall = int(ats * 0.4 + read * 0.3 + imp * 0.3)
    scores = ScoreBreakdown(ats_score=ats, readability_score=read,
                            impact_score=imp, overall_score=overall)
    sug = generate_suggestions(text, scores, skills, missing)
    qns = generate_questions(text, skills)
    save_to_history(filename, scores, sug, skills)
    return AnalysisResult(suggestions=sug, interview_questions=qns,
                          scores=scores, detected_skills=skills,
                          missing_sections=missing, resume_text=text[:6000])

# ---------------------------------------------------------------------------
# JD Matching
# ---------------------------------------------------------------------------
SKILL_RESOURCES = {
    "Python": "https://docs.python.org/3/tutorial/",
    "JavaScript": "https://javascript.info/",
    "React": "https://react.dev/learn",
    "Node.js": "https://nodejs.org/en/learn",
    "SQL": "https://www.w3schools.com/sql/",
    "AWS": "https://aws.amazon.com/training/",
    "Docker": "https://docs.docker.com/get-started/",
    "Kubernetes": "https://kubernetes.io/docs/tutorials/",
    "Machine Learning": "https://www.coursera.org/learn/machine-learning",
    "Deep Learning": "https://www.deeplearningbook.org/",
    "TensorFlow": "https://www.tensorflow.org/tutorials",
    "PyTorch": "https://pytorch.org/tutorials/",
    "Git": "https://git-scm.com/doc",
    "TypeScript": "https://www.typescriptlang.org/docs/",
    "Data Science": "https://www.kaggle.com/learn",
    "MongoDB": "https://university.mongodb.com/",
    "Azure": "https://learn.microsoft.com/en-us/training/azure/",
    "GCP": "https://cloud.google.com/training",
    "Agile": "https://www.scrum.org/resources",
    "Django": "https://docs.djangoproject.com/",
    "Flask": "https://flask.palletsprojects.com/",
    "CI/CD": "https://www.atlassian.com/continuous-delivery",
}

STOP = {"the","a","an","is","are","was","were","be","been","have","has","had",
        "do","does","did","will","would","could","should","may","might","can",
        "and","but","or","nor","not","so","yet","both","either","neither",
        "all","any","few","more","most","other","some","no","only","own",
        "same","than","too","very","just","now","our","we","you","your",
        "they","them","their","its","for","with","about","between","through",
        "during","before","after","to","from","up","down","in","out","on",
        "off","over","under","again","then","once","here","there","when",
        "where","why","how","what","which","who","whom","this","that","these",
        "those","am","at","by","of","if","into","it","as","i","me","my","he",
        "him","his","she","her","must","also","work","working","experience",
        "role","job","position","required","requirements","looking","able",
        "strong","knowledge","understanding","skills","years","using","used",
        "well","new","etc","such"}

def match_jd_local(resume_text: str, jd: str) -> JDMatchResult:
    jd_skills = set(detect_skills(jd))
    res_skills = set(detect_skills(resume_text))
    jd_words = set(re.findall(r'\b[a-z]{4,}\b', jd.lower())) - STOP
    res_lower = resume_text.lower()
    matched_words = [w for w in jd_words if w in res_lower]
    word_pct = (len(matched_words) / max(len(jd_words), 1)) * 50
    matched_sk = jd_skills & res_skills
    missing_sk = jd_skills - res_skills
    skill_pct = (len(matched_sk) / max(len(jd_skills), 1)) * 50
    score = max(0, min(100, int(word_pct + skill_pct)))
    gaps = []
    for sk in missing_sk:
        g = {"skill": sk, "suggestion": f"Learn {sk} to match this role."}
        if sk in SKILL_RESOURCES: g["resource"] = SKILL_RESOURCES[sk]
        gaps.append(g)
    return JDMatchResult(match_score=score,
                         matched_keywords=sorted(matched_sk),
                         missing_keywords=sorted(missing_sk),
                         skill_gaps=gaps)

# ---------------------------------------------------------------------------
# Mock Interview Evaluation
# ---------------------------------------------------------------------------
def evaluate_answer_local(question: str, answer: str, ctx: str = "") -> InterviewEvalResult:
    a_low, wc = answer.lower(), len(answer.split())
    sents = [s.strip() for s in re.split(r'[.!?]+', answer) if s.strip()]

    # Confidence
    conf = 50
    if wc >= 50: conf += 15
    if wc >= 100: conf += 10
    if any(w in a_low for w in ["for example","specifically","in particular"]): conf += 10
    if any(w in a_low for w in ["i achieved","i led","i built","i developed"]): conf += 10
    if re.search(r'\d+', answer): conf += 5
    conf = min(100, conf)

    # Clarity
    clar = 60
    if len(sents) >= 3: clar += 10
    avg_sl = sum(len(s.split()) for s in sents) / max(len(sents), 1)
    if 10 <= avg_sl <= 20: clar += 15
    elif avg_sl > 30: clar -= 10
    if wc < 20: clar -= 20
    clar = max(0, min(100, clar))

    # Relevance
    q_words = set(re.findall(r'\b\w+\b', question.lower())) - STOP
    a_words = set(re.findall(r'\b\w+\b', a_low))
    rel = min(100, 40 + len(q_words & a_words) * 10)

    overall = int(conf * 0.35 + clar * 0.35 + rel * 0.3)

    parts = []
    if conf >= 70: parts.append("You show good confidence with specific details.")
    else: parts.append("Be more specific — include concrete examples.")
    if clar >= 70: parts.append("Well-structured and easy to follow.")
    else: parts.append("Use the STAR method (Situation, Task, Action, Result).")
    if rel >= 70: parts.append("Answer directly addresses the question.")
    else: parts.append("Focus more directly on what's being asked.")

    tips = []
    if wc < 50: tips.append("Expand with more detail — aim for 50–150 words.")
    if not any(w in a_low for w in ["example","instance","project","when i"]): tips.append("Include a specific example or anecdote.")
    if not re.search(r'\d+', answer): tips.append("Add numbers where possible (e.g., 'improved by 30%').")
    if conf < 70: tips.append("Structure with STAR: Situation → Task → Action → Result.")
    tips.append("Practice out loud for natural delivery.")

    return InterviewEvalResult(confidence_score=conf, clarity_score=clar,
                               relevance_score=rel, overall_score=overall,
                               feedback=" ".join(parts), improvement_tips=tips[:4])

# ---------------------------------------------------------------------------
# History Persistence
# ---------------------------------------------------------------------------
def load_history() -> list:
    try:
        with open(HISTORY_FILE, "r") as f: return json.load(f)
    except (FileNotFoundError, json.JSONDecodeError): return []

def save_to_history(filename, scores, suggestions, skills):
    hist = load_history()
    hist.append({
        "id": len(hist) + 1,
        "filename": filename,
        "date": datetime.now().isoformat(),
        "scores": scores.model_dump() if hasattr(scores, 'model_dump') else scores.dict(),
        "suggestion_count": len(suggestions),
        "skills_detected": len(skills),
        "top_skills": skills[:5],
    })
    with open(HISTORY_FILE, "w") as f: json.dump(hist, f, indent=2)

# ---------------------------------------------------------------------------
# PDF Report Generation
# ---------------------------------------------------------------------------
def generate_pdf_report(data: dict) -> bytes:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    # Title
    pdf.set_font("Helvetica", "B", 22)
    pdf.set_text_color(59, 130, 246)
    pdf.cell(0, 15, "AI Resume Coach Report", ln=True, align="C")
    pdf.set_text_color(100, 116, 139)
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(0, 8, f"Generated {datetime.now().strftime('%B %d, %Y at %I:%M %p')}", ln=True, align="C")
    pdf.ln(10)

    scores = data.get("scores", {})
    pdf.set_text_color(30, 41, 59)
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "Score Breakdown", ln=True)
    pdf.set_draw_color(59, 130, 246)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y()); pdf.ln(5)
    pdf.set_font("Helvetica", "", 12)
    for label, key in [("Overall", "overall_score"), ("ATS", "ats_score"),
                       ("Readability", "readability_score"), ("Impact", "impact_score")]:
        v = scores.get(key, 0)
        pdf.cell(50, 8, f"{label}:")
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(0, 8, f"{v} / 100", ln=True)
        pdf.set_font("Helvetica", "", 12)
    pdf.ln(5)

    skills = data.get("detected_skills", [])
    if skills:
        pdf.set_font("Helvetica", "B", 16)
        pdf.cell(0, 10, "Detected Skills", ln=True)
        pdf.set_draw_color(139, 92, 246)
        pdf.line(10, pdf.get_y(), 200, pdf.get_y()); pdf.ln(5)
        pdf.set_font("Helvetica", "", 11)
        pdf.multi_cell(0, 7, ", ".join(skills)); pdf.ln(3)

    sug = data.get("suggestions", [])
    if sug:
        pdf.set_font("Helvetica", "B", 16)
        pdf.cell(0, 10, "Improvement Suggestions", ln=True)
        pdf.set_draw_color(59, 130, 246)
        pdf.line(10, pdf.get_y(), 200, pdf.get_y()); pdf.ln(5)
        pdf.set_font("Helvetica", "", 11)
        for i, s in enumerate(sug, 1):
            clean = re.sub(r'[^\x00-\x7F]+', '', s).strip()
            pdf.multi_cell(0, 7, f"{i}. {clean}"); pdf.ln(2)

    qs = data.get("interview_questions", [])
    if qs:
        pdf.set_font("Helvetica", "B", 16)
        pdf.cell(0, 10, "Interview Questions", ln=True)
        pdf.set_draw_color(139, 92, 246)
        pdf.line(10, pdf.get_y(), 200, pdf.get_y()); pdf.ln(5)
        pdf.set_font("Helvetica", "", 11)
        for i, q in enumerate(qs, 1):
            pdf.multi_cell(0, 7, f"Q{i}. {q}"); pdf.ln(2)

    jd = data.get("jd_match")
    if jd:
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 16)
        pdf.cell(0, 10, f"Job Description Match: {jd.get('match_score', 0)}%", ln=True)
        pdf.set_draw_color(59, 130, 246)
        pdf.line(10, pdf.get_y(), 200, pdf.get_y()); pdf.ln(5)
        mk = jd.get("matched_keywords", [])
        ms = jd.get("missing_keywords", [])
        if mk:
            pdf.set_font("Helvetica", "B", 12); pdf.cell(0, 8, "Matched:", ln=True)
            pdf.set_font("Helvetica", "", 11); pdf.multi_cell(0, 7, ", ".join(mk)); pdf.ln(3)
        if ms:
            pdf.set_font("Helvetica", "B", 12); pdf.cell(0, 8, "Missing:", ln=True)
            pdf.set_font("Helvetica", "", 11); pdf.multi_cell(0, 7, ", ".join(ms))
        gaps = jd.get("skill_gaps", [])
        if gaps:
            pdf.ln(5)
            pdf.set_font("Helvetica", "B", 12); pdf.cell(0, 8, "Skill Gaps & Resources:", ln=True)
            pdf.set_font("Helvetica", "", 11)
            for g in gaps:
                line = f"- {g['skill']}: {g.get('suggestion','')}"
                if g.get('resource'): line += f"  ({g['resource']})"
                pdf.multi_cell(0, 7, line); pdf.ln(1)

    return bytes(pdf.output())

# ========================== ENDPOINTS ==========================

@app.post("/analyze", response_model=AnalysisResult)
async def analyze_resume(file: UploadFile = File(...)):
    """Upload PDF/DOCX resume and receive full scored analysis."""
    if not file.filename.lower().endswith(('.pdf', '.docx')):
        raise HTTPException(400, "Only PDF or DOCX files are allowed.")
    
    contents = await file.read()
    if len(contents) > 10 * 1024 * 1024:
        raise HTTPException(400, "File size exceeds 10MB limit.")
        
    try:
        text = extract_text(contents, file.filename)
    except Exception as e:
        raise HTTPException(500, f"Extraction error: {e}")
    if not text.strip():
        raise HTTPException(400, "Could not extract text from the document.")
    return local_analyze(text, file.filename)

@app.post("/match-jd", response_model=JDMatchResult)
async def match_job_description(req: JDMatchRequest):
    """Match resume text against a job description."""
    if not req.resume_text.strip() or not req.job_description.strip():
        raise HTTPException(400, "Both resume text and JD are required.")
    return match_jd_local(req.resume_text, req.job_description)

@app.post("/interview/evaluate", response_model=InterviewEvalResult)
async def evaluate_interview(req: InterviewEvalRequest):
    """Evaluate a mock interview answer."""
    if not req.answer.strip():
        raise HTTPException(400, "Answer cannot be empty.")
    return evaluate_answer_local(req.question, req.answer, req.resume_context)

@app.get("/history")
async def get_history():
    """Return all past analysis entries."""
    return load_history()

@app.post("/report")
async def download_report(req: ReportRequest):
    """Generate and download a PDF report of the analysis."""
    try:
        pdf_bytes = generate_pdf_report(req.model_dump() if hasattr(req,'model_dump') else req.dict())
        return StreamingResponse(
            io.BytesIO(pdf_bytes),
            media_type="application/pdf",
            headers={"Content-Disposition": "attachment; filename=resume_coach_report.pdf"},
        )
    except Exception as e:
        raise HTTPException(500, f"Report generation failed: {e}")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)
